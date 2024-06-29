import cv2
import numpy as np
import easyocr
import os
import warnings
import re
import pandas as pd
from difflib import SequenceMatcher
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage

# Suppress FutureWarnings from PyTorch
warnings.filterwarnings("ignore", category=FutureWarning)
reader = easyocr.Reader(['en', 'pl'])


# --------------------- Video & ROI Processing Functions --------------------- #

def initialize_video(video_path):
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        raise ValueError("Error: Could not open video.")
    fps = cap.get(cv2.CAP_PROP_FPS)
    if fps == 0:
        cap.release()
        raise ValueError("Error: FPS value is 0. Cannot process video.")
    return cap, fps

def resize_frame(frame, new_width, new_height):
    return cv2.resize(frame, (new_width, new_height))

def extract_roi(frame, roi_box_position, roi_box_size):
    return frame[
        roi_box_position[1]:roi_box_position[1] + roi_box_size[1],
        roi_box_position[0]:roi_box_position[0] + roi_box_size[0]
    ]

def detect_circles(roi_img):
    blurred = cv2.GaussianBlur(roi_img, (9, 9), 2)
    channels = cv2.split(blurred)
    circles_list = []
    for channel in channels:
        circles = cv2.HoughCircles(
            channel,
            cv2.HOUGH_GRADIENT,
            dp=1.2,
            minDist=100,
            param1=100,
            param2=30,
            minRadius=10,
            maxRadius=30
        )
        if circles is not None:
            circles = np.uint16(np.around(circles))
            circles_list.extend(circles[0, :])
    return circles_list

def Crop_Frame_TO_AconutName(circles_list, roi_box_position, resized_frame, display_width, display_height):
    # Keep only the last circle regarding the y-axis
    largest_circle = max(circles_list, key=lambda c: c[1])
    x, y, r = largest_circle
    center = (x + roi_box_position[0], y + roi_box_position[1])
    new_roi_width, new_roi_height = 380, 40
    new_roi_x = center[0] + r
    new_roi_y = center[1] - new_roi_height // 2
    new_roi_x = min(new_roi_x, display_width - new_roi_width)
    new_roi_y = max(min(new_roi_y, display_height - new_roi_height), 0)
    acconut_name_img = resized_frame[
        new_roi_y:new_roi_y + new_roi_height,
        new_roi_x:new_roi_x + new_roi_width
    ]
    return acconut_name_img


# --------------------- OCR & Filtering Functions --------------------- #

def is_valid_instagram_username(username):
    """Check if the Instagram username is valid."""
    if re.search(r'[@?#$%^&*(),"{}|<>`=\[\]\\;\'/]', username):
        return False
    if username.endswith('.') or username.startswith('.'):
        return False
    if any(substring in username for substring in ['Follo', 'ollow', 'Fol']):
        return False
    return True

def filter_non_usernames(usernames):
    """Filter out invalid Instagram usernames."""
    return [username for username in usernames if is_valid_instagram_username(username)]

def is_number(text):
    """Check if the text is a number."""
    try:
        float(text.replace(',', ''))
        return True
    except ValueError:
        return False

def OCR(frame, roi=None):
    """Perform OCR on the given frame."""
    if roi:
        x, y, w, h = roi
        frame = frame[y:y+h, x:x+w]
    result = reader.readtext(frame, paragraph=True, x_ths=1.0, y_ths=0.1, text_threshold=0.8)
    return [item[1] for item in result if not is_number(item[1])]

def OCR_count(OCR_results, dic, min_diff=3):
    """Count occurrences of OCR results and return the top k most repeated account names."""
    for result in OCR_results:
        dic[result] = dic.get(result, 0) + 1
        for key in dic.keys():
            if key not in OCR_results and (result in key or key in result):
                dic[key] += 1

    sorted_results = sorted(dic.items(), key=lambda item: item[1], reverse=True)
    top_k_results = []
    for i in range(len(sorted_results)):
        if i == 0 or (sorted_results[i-1][1] - sorted_results[i][1] <= min_diff):
            top_k_results.append(sorted_results[i])
        else:
            break

    top_k_account_names = [item[0] for item in top_k_results]
    top_k_counts = [item[1] for item in top_k_results]
    return dic, top_k_counts, top_k_account_names

def split_and_strip_usernames(usernames):
    """Split usernames if they contain 'and' or 'with' and strip extra spaces."""
    split_usernames = []
    for username in usernames:
        parts = re.split(r'\band\b|\bwith\b', username)
        for part in parts:
            stripped_part = part.strip()
            if stripped_part:
                split_usernames.append(stripped_part)
    return split_usernames

def OCR_pipeline(dic, frame):
    """Process a frame through the OCR pipeline."""
    OCR_results = OCR(frame)
    usernames = filter_non_usernames(OCR_results)
    usernames = split_and_strip_usernames(usernames)
    dic, top_k_counts, top_k_account_names = OCR_count(usernames, dic)
    return top_k_counts, top_k_account_names, usernames


# --------------------- Video Processing Functions --------------------- #

def process_video(video_path, checks_per_sec=3, number_of_sec=3):
    cap, fps = initialize_video(video_path)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration_sec = total_frames / fps
    if duration_sec < 3:
        cap.release()
        return None

    skip_frames = int(fps / checks_per_sec)
    display_width, display_height = 540, 960
    roi_box_size = (69, 250)
    roi_box_position = (10, display_height - 400)
    dic = {}
    frames_to_process = list(range(0, total_frames, skip_frames))

    for frame_number in frames_to_process:
        cap.set(cv2.CAP_PROP_POS_FRAMES, frame_number)
        ret, frame = cap.read()
        if not ret:
            break

        resized_frame = resize_frame(frame, display_width, display_height)
        roi_img = extract_roi(resized_frame, roi_box_position, roi_box_size)
        circles_list = detect_circles(roi_img)
        if circles_list:
            account_name_frame = Crop_Frame_TO_AconutName(
                circles_list, roi_box_position, resized_frame, display_width, display_height
            )
            top_k_counts, top_k_account_names, usernames = OCR_pipeline(dic, account_name_frame)
            if top_k_counts and top_k_counts[0] > number_of_sec * checks_per_sec:
                dic = {}
                cap.release()
                return top_k_account_names
    cap.release()
    if circles_list:
        return top_k_account_names

def process_video_reverse(video_path, checks_per_sec=3, number_of_sec=3):
    cap, fps = initialize_video(video_path)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration_sec = total_frames / fps
    if duration_sec < 3:
        cap.release()
        return None

    skip_frames = int(fps / checks_per_sec)
    display_width, display_height = 540, 960
    roi_box_size = (69, 250)
    roi_box_position = (10, display_height - 400)
    dic = {}
    frames_to_process = list(range(total_frames - 1, -1, -skip_frames))
    top_k_account_names = None

    for frame_number in frames_to_process:
        cap.set(cv2.CAP_PROP_POS_FRAMES, frame_number)
        ret, frame = cap.read()
        if not ret:
            break

        resized_frame = resize_frame(frame, display_width, display_height)
        roi_img = extract_roi(resized_frame, roi_box_position, roi_box_size)
        circles_list = detect_circles(roi_img)
        if circles_list:
            account_name_frame = Crop_Frame_TO_AconutName(
                circles_list, roi_box_position, resized_frame, display_width, display_height
            )
            top_k_counts, top_k_account_names, usernames = OCR_pipeline(dic, account_name_frame)
            if top_k_counts and top_k_counts[0] > number_of_sec * checks_per_sec:
                dic = {}
                cap.release()
                return top_k_account_names
    cap.release()
    if top_k_account_names:
        return top_k_account_names
    return None

def similar(a, b):
    return SequenceMatcher(None, a, b).ratio()

def find_similar_names(list1, list2, threshold=0.8):
    similar_names = []
    for name1 in list1:
        for name2 in list2:
            similarity = similar(name1, name2)
            if similarity >= threshold:
                similar_names.append((name1, name2, similarity))
    return similar_names

def extract_part_number(filename):
    match = re.search(r'part_(\d+)', filename)
    return int(match.group(1)) if match else float('inf')

def instagram_wrapper(video_path):
    # Process forward and reverse
    top_k_account_names_forward = process_video(video_path, checks_per_sec=20, number_of_sec=2)
    top_k_account_names_reverse = process_video_reverse(video_path, checks_per_sec=20, number_of_sec=2)
    if top_k_account_names_forward is None:
        union_names = top_k_account_names_reverse
    elif top_k_account_names_reverse is None:
        union_names = top_k_account_names_forward
    else:
        union_names = list(dict.fromkeys(top_k_account_names_forward + top_k_account_names_reverse))
    more_than_one_username = 0 if (top_k_account_names_forward and top_k_account_names_reverse) else 1
    return [video_path, top_k_account_names_forward, top_k_account_names_reverse, union_names, more_than_one_username]


# --------------------- Screenshot Capture Helper --------------------- #

def get_roi_screenshot(cap, frame_idx, display_width, display_height, roi_box_position, roi_box_size):
    cap.set(cv2.CAP_PROP_POS_FRAMES, frame_idx)
    ret, frame = cap.read()
    if not ret:
        return None
    resized_frame = resize_frame(frame, display_width, display_height)
    roi_img = extract_roi(resized_frame, roi_box_position, roi_box_size)
    circles_list = detect_circles(roi_img)
    if circles_list:
        account_name_frame = Crop_Frame_TO_AconutName(
            circles_list, roi_box_position, resized_frame, display_width, display_height
        )
        return account_name_frame
    return roi_img


# --------------------- Main Processing & Excel Export --------------------- #

def main(folder_path, excel_output):
    results = []
    video_files = sorted(
        [f for f in os.listdir(folder_path) if f.endswith('.mp4')],
        key=extract_part_number
    )
    screenshot_folder = "screenshots"
    if not os.path.exists(screenshot_folder):
        os.makedirs(screenshot_folder)
        
    # Constants for frame processing:
    display_width, display_height = 540, 960
    roi_box_position = (10, display_height - 400)
    roi_box_size = (69, 250)
    
    for video_file in video_files:
        video_path = os.path.join(folder_path, video_file)
        result = instagram_wrapper(video_path)  # [video_path, forward, reverse, union, flag]
        
        # Open video to capture the screenshot at the 3-second mark.
        cap, fps = initialize_video(video_path)
        total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        duration_sec = total_frames / fps
        
        screenshot_path = ""
        if duration_sec >= 3:
            frame_idx = int(fps * 3)
            screenshot_img = get_roi_screenshot(cap, frame_idx, display_width, display_height, roi_box_position, roi_box_size)
            if screenshot_img is not None:
                screenshot_path = os.path.join(screenshot_folder, f"{os.path.splitext(video_file)[0]}_3sec.png")
                cv2.imwrite(screenshot_path, screenshot_img)
        cap.release()
        
        # Append result row: video path, forward result, reverse result, union names, flag, and screenshot path.
        results.append([
            video_path,
            ", ".join(result[1]) if result[1] is not None else "",
            ", ".join(result[2]) if result[2] is not None else "",
            ", ".join(result[3]) if result[3] is not None else "",
            result[4],
            screenshot_path
        ])
    
    # Create an Excel workbook and add the screenshot image.
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = ["Video Path", "Forward Result", "Reverse Result", "Union Names", "More Than One Username", "Screenshot 3sec"]
    ws.append(headers)
    
    row_num = 2
    for row in results:
        ws.cell(row=row_num, column=1, value=row[0])
        ws.cell(row=row_num, column=2, value=row[1])
        ws.cell(row=row_num, column=3, value=row[2])
        ws.cell(row=row_num, column=4, value=row[3])
        ws.cell(row=row_num, column=5, value=row[4])
        # Insert image if the screenshot exists.
        if row[5] and os.path.exists(row[5]):
            img = ExcelImage(row[5])
            cell_ref = ws.cell(row=row_num, column=6).coordinate
            ws.add_image(img, cell_ref)
        row_num += 1

    wb.save(excel_output)

if __name__ == "__main__":
    # Update these paths as needed for your environment.
    folder_path = r"video_folder"
    excel_output = "results.xlsx"
    
    try:
        main(folder_path, excel_output)
    except KeyboardInterrupt:
        pass
