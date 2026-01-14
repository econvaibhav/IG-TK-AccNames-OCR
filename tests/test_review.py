"""Corrections must survive reload, preserve raw data and reject stale saves."""

import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from ig_tk_ocr.core import Config
from ig_tk_ocr.pipeline import screenshots
from ig_tk_ocr.review import Conflict, ReviewStore, make_server, row_id


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.row = {"path":"sample.mp4", "status":"agreement", "union_names":["name-with-dot"], "needs_review":True}
        self.raw = json.dumps(self.row) + "\n"
        (self.folder / "details.jsonl").write_text(self.raw)
        (self.folder / "results.csv").write_text("original\n")
        self.store = ReviewStore(self.folder)
        self.payload = {"id":row_id(self.row),"revision":0,"names":["name.with.dot"], "reviewed":True, "unreadable":False, "notes":"=literal note"}

    def test_save_reload_and_spreadsheets_preserve_original(self):
        errors = self.store.save(self.payload)
        self.assertEqual(errors, [])
        reloaded = ReviewStore(self.folder)
        self.assertEqual(reloaded.state["entries"][self.payload["id"]]["names"], ["name.with.dot"])
        self.assertEqual((self.folder / "details.jsonl").read_text(), self.raw)
        self.assertEqual((self.folder / "results.csv").read_text(), "original\n")
        self.assertFalse((self.folder / "reviewed.csv").exists())
        from openpyxl import load_workbook
        book = load_workbook(self.folder / "reviewed.xlsx")
        sheet = book.active
        headers = [c.value for c in sheet[1]]
        cell = sheet.cell(2, headers.index("review_notes")+1)
        self.assertEqual(cell.value, "=literal note")
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(sheet.cell(2, headers.index("final_names")+1).value, '["name.with.dot"]')
        book.close()

    def test_stale_revision_does_not_overwrite(self):
        self.store.save(self.payload)
        with self.assertRaises(Conflict):
            self.store.save({**self.payload, "names":"different"})
        self.assertEqual(self.store.state["revision"], 1)

    def test_unchanged_non_ascii_and_display_names_can_be_confirmed(self):
        for name in ("България Новини", "Müller Politik", "Łódź, Polska", "a-name", "very.long.display.name.over.thirty.characters"):
            self.store.by_id[self.payload["id"]]["union_names"] = [name]
            self.store.save({**self.payload, "revision":self.store.state["revision"], "names":[name]})
            entry = self.store.state["entries"][self.payload["id"]]
            self.assertEqual(entry["decision"], "confirmed")
            self.assertEqual(entry["names"], [name])

    def test_changed_name_is_automatically_corrected(self):
        self.store.save(self.payload)
        self.assertEqual(self.store.state["entries"][self.payload["id"]]["decision"], "corrected")

    def test_empty_review_and_control_characters_rejected(self):
        for changes in ({"names":[]}, {"names":["bad\x00name"]}, {"names":["x"*201]}, {"unreadable":True}):
            with self.assertRaises(ValueError):
                self.store.save({**self.payload, **changes})
        self.assertFalse(self.store.path.exists())

    def test_unchecked_review_keeps_draft_and_original_final_names(self):
        self.store.save({**self.payload,"reviewed":False})
        entry = ReviewStore(self.folder).state["entries"][self.payload["id"]]
        self.assertEqual(entry["names"], ["name.with.dot"])
        self.assertEqual(entry["decision"], "unreviewed")
        self.assertEqual(json.loads(self.store.reviewed_rows()[0]["final_names"]),["name-with-dot"])

    def test_old_saved_review_and_old_page_api_are_compatible(self):
        self.store.save({"id":self.payload["id"],"revision":0,"names":"България Новини","decision":"confirmed"})
        reloaded=ReviewStore(self.folder)
        self.assertEqual(reloaded.state["entries"][self.payload["id"]]["names"],["България Новини"])
        self.assertEqual(reloaded.state["entries"][self.payload["id"]]["decision"],"corrected")

    def test_unreadable_clears_final_names_and_remains_unresolved(self):
        self.store.save({**self.payload, "reviewed":False, "unreadable":True})
        row = self.store.reviewed_rows()[0]
        self.assertEqual(row["final_names"], "[]")
        self.assertTrue(row["needs_manual_review"])

    def test_excel_lock_keeps_correction_durable(self):
        with patch("ig_tk_ocr.excel.save_excel", side_effect=PermissionError("locked")):
            errors = self.store.save(self.payload)
        self.assertIn("reviewed.xlsx", errors[0])
        self.assertEqual(ReviewStore(self.folder).state["revision"], 1)
        self.assertEqual(self.store.export(), [])

    def test_local_http_saves_and_rejects_foreign_requests(self):
        server = make_server(self.store)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            url = f"http://127.0.0.1:{server.server_port}"
            with urlopen(url + "/api/state") as response:
                state = json.load(response)
            body = json.dumps(self.payload).encode()
            headers = {"Content-Type":"application/json", "X-Review-Token":state["token"]}
            with urlopen(Request(url + "/api/save", body, headers)) as response:
                self.assertEqual(json.load(response)["revision"], 1)
            for route, hdr in (("/api/save",{**headers,"Origin":"https://foreign.example"}),
                               ("/api/save",{"Content-Type":"application/json"})):
                with self.assertRaises(HTTPError) as error:
                    urlopen(Request(url + route, body, hdr))
                self.assertEqual(error.exception.code, 403)
            with self.assertRaises(HTTPError):
                urlopen(url + "/screenshots/../../details.jsonl")
            with self.assertRaises(HTTPError) as csv_error:
                urlopen(url + "/reviewed.csv")
            self.assertEqual(csv_error.exception.code, 404)
        finally:
            server.shutdown(); server.server_close(); thread.join()


class FrameTests(unittest.TestCase):
    def test_quartiles_legacy_points_and_short_clip_bounds(self):
        import numpy as np
        class Capture:
            indices = []
            def set(self, key, index): self.indices.append(index)
            def read(self): return True, np.zeros((96,54,3),dtype=np.uint8)
        for count in (4,100):
            cap = Capture(); cap.indices = []
            with tempfile.TemporaryDirectory() as tmp, patch("ig_tk_ocr.pipeline.crop_frame",return_value=(np.zeros((96,54,3),dtype=np.uint8),None,None)):
                shots = screenshots(cap,10,count,Config(),Path(tmp))
                self.assertEqual(len(shots), 6)
                self.assertTrue(all(0 <= index < count for index in cap.indices))
                quartiles = [s for s in shots if "percent" in s["label"]]
                self.assertEqual([s["frame"] for s in quartiles], [round(count*.25),round(count*.5),round(count*.75)])
                self.assertTrue(all(Path(s["full_path"]).is_file() for s in quartiles))
                if count == 4:
                    self.assertEqual(shots[0]["error"], "outside_clip")


if __name__ == "__main__":
    unittest.main()
