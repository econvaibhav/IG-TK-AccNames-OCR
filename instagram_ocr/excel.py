"""Optional portable Excel exporter, retaining the original screenshot workflow."""

from pathlib import Path

from .files import csv_row


def save_excel(results, path):
    # Optional dependency: basic OCR and CSV/HTML export do not need openpyxl.
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Results"
    headers = list(csv_row({"path": "", "status": ""})) + ["Screenshot at 3 s"]
    sheet.append(headers)
    sheet.freeze_panes = "B2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174F60")
    for row_index, result in enumerate(results, 2):
        sheet.append(list(csv_row(result).values()))
        for cell in sheet[row_index]:
            # Treat OCR strings and paths as literal text, including leading '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row_index].height = 70
        shot = next((s for s in result.get("screenshots", []) if s["label"] == "3sec" and "path" in s), None)
        if shot and Path(shot["path"]).is_file():
            image = Image(shot["path"])
            scale = min(1, 380 / image.width, 80 / image.height)
            image.width *= scale
            image.height *= scale
            sheet.add_image(image, f"K{row_index}")
    for column in "ABCD EFGHIJ".replace(" ", ""):
        sheet.column_dimensions[column].width = 24 if column != "A" else 42
    sheet.column_dimensions["K"].width = 55
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)

